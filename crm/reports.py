import calendar
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP, localcontext
from zoneinfo import ZoneInfo

from django.contrib.auth import get_user_model
from django.db.models import Count, Q
from django.utils import timezone

from .catalog import paginate
from .models import (
    Brand,
    CommissionAdjustment,
    CommissionReview,
    InventoryBalance,
    InventoryHold,
    PaymentAccount,
    PaymentRecord,
    SalesOrder,
    SalesOrderLine,
    SKU,
    Warehouse,
)
from .operations import OperationError, identifier, inventory_rows
from .sales import sales_order_queryset_for


REPORT_TIME_ZONE = ZoneInfo("Asia/Shanghai")
ZERO = Decimal("0.000")
EXCEL_MAX_ROWS = 1_048_576

REPORTS = {
    "inventory": {"title": "库存报表", "permission": "view_inventorybalance"},
    "sales": {"title": "销售报表", "permission": "view_salesorder"},
    "payments": {"title": "收款核对报表", "permission": "view_salesorder"},
    "commissions": {"title": "提成报表", "permission": "view_salesorder"},
}

WORKBENCH_REPORTS = {
    "sales": {"customer_ranking", "day_over_day"},
    "inventory": {"ip_language_stock_count"},
}

ORDER_FILTER_LABELS = {
    "brand_id": "IP",
    "sku_id": "SKU",
    "warehouse_id": "仓库",
    "salesperson_id": "销售",
    "customer_id": "客户",
    "channel_id": "渠道",
    "province": "收货省份",
    "city": "收货城市",
}


def _now_local():
    return timezone.now().astimezone(REPORT_TIME_ZONE)


def _aware_midnight(day):
    return datetime.combine(day, time.min, tzinfo=REPORT_TIME_ZONE)


def _date_value(value, field):
    if not isinstance(value, str) or not value:
        raise OperationError("日期格式无效。", fields={field: "请按YYYY-MM-DD填写日期。"})
    try:
        parsed = date.fromisoformat(value)
    except ValueError:
        raise OperationError("日期格式无效。", fields={field: "请按YYYY-MM-DD填写日期。"}) from None
    if parsed.isoformat() != value:
        raise OperationError("日期格式无效。", fields={field: "请按YYYY-MM-DD填写日期。"})
    return parsed


def _months_before(day, months):
    month_index = day.year * 12 + day.month - 1 - months
    year, month_zero = divmod(month_index, 12)
    month = month_zero + 1
    return date(year, month, min(day.day, calendar.monthrange(year, month)[1]))


def _date_criteria(query, *, today=None):
    period = query.get("period", "")
    start_text = query.get("start_date", "")
    end_text = query.get("end_date", "")
    if period in (None, "") and (start_text or end_text):
        period = "custom"
    period = period or "all"
    valid = {"1m": 1, "3m": 3, "6m": 6, "12m": 12, "all": None, "custom": None}
    if period not in valid:
        raise OperationError("统计周期无效。", fields={"period": "请选择近1个月、近3个月、近6个月、近12个月、全部或自定义。"})
    if period == "custom":
        if not start_text or not end_text:
            raise OperationError("自定义周期需要起止日期。", fields={"start_date": "请同时填写开始和结束日期。", "end_date": "请同时填写开始和结束日期。"})
        start_day = _date_value(start_text, "start_date")
        end_day = _date_value(end_text, "end_date")
        if start_day > end_day:
            raise OperationError("开始日期不能晚于结束日期。", fields={"start_date": "请检查起止日期。"})
    elif valid[period] is None:
        if start_text or end_text:
            raise OperationError("全部周期不能同时填写自定义日期。", fields={"period": "请选择自定义周期后再填写日期。"})
        start_day = end_day = None
    else:
        if start_text or end_text:
            raise OperationError("滚动周期不能同时填写自定义日期。", fields={"period": "请选择自定义周期后再填写日期。"})
        end_day = today or _now_local().date()
        start_day = _months_before(end_day, valid[period])
    try:
        end_before = _aware_midnight(end_day + timedelta(days=1)) if end_day else None
    except OverflowError:
        raise OperationError(
            "结束日期超出可筛选范围。",
            fields={"end_date": "结束日期无法计算次日边界。"},
        ) from None
    return {
        "period": period,
        "start_date": start_day.isoformat() if start_day else None,
        "end_date": end_day.isoformat() if end_day else None,
        "start_at": _aware_midnight(start_day) if start_day else None,
        "end_before": end_before,
        "timezone": "Asia/Shanghai",
        "period_label": (
            f"{start_day.isoformat()} 至 {end_day.isoformat()}（含首尾，北京时间）"
            if start_day else "全部日期"
        ),
    }


def _filter_identifier(query, field):
    value = query.get(field, "")
    return identifier(value, field) if value not in (None, "") else None


def _reject_workbench_filters(query, fields):
    for field in fields:
        if query.get(field) not in (None, ""):
            raise OperationError("工作台统计不能与此筛选组合。", fields={field: "请移除该筛选。"})


def _order_queryset(actor):
    return sales_order_queryset_for(actor).select_related("salesperson", "customer", "channel")


def _order_filters(orders, query):
    for field, model_field in (
        ("salesperson_id", "salesperson_id"),
        ("customer_id", "customer_id"),
        ("channel_id", "channel_id"),
    ):
        value = _filter_identifier(query, field)
        if value is not None:
            orders = orders.filter(**{model_field: value})
    for field in ("province", "city"):
        value = query.get(field, "").strip()
        if value:
            orders = orders.filter(**{f"{field}_snapshot": value})
    line_filters = {}
    for field, model_field in (
        ("brand_id", "sku__product__brand_id"),
        ("sku_id", "sku_id"),
        ("warehouse_id", "warehouse_id"),
    ):
        value = _filter_identifier(query, field)
        if value is not None:
            line_filters[model_field] = value
    if line_filters:
        matching_lines = SalesOrderLine.objects.filter(order_id__in=orders.values("pk"), active=True, **line_filters)
        orders = orders.filter(pk__in=matching_lines.values("order_id"))
    return orders


def _filter_summary(query, metadata):
    parts = [f"周期：{metadata['period_label']}"]
    for field, label in ORDER_FILTER_LABELS.items():
        value = query.get(field, "").strip()
        if value:
            parts.append(f"{label}={value}")
    if query.get("scope") == "unshipped":
        parts.append("范围=未交付订单（不按交付日期筛选）")
    for field, label in (("status", "核对状态"), ("account_id", "收款账户"), ("paused", "暂停状态"), ("commission_status", "提成状态")):
        value = query.get(field, "").strip()
        if value:
            parts.append(f"{label}={value}")
    return "；".join(parts)


def _with_query_metadata(report, query, *, basis, date_filter=True, payment_pending=False, today=None):
    metadata = {
        "timezone": "Asia/Shanghai",
        "queried_at": _now_local().isoformat(timespec="seconds"),
        "date_basis": basis,
        "date_filter_available": date_filter,
        "payment_date_basis_pending": payment_pending,
    }
    if date_filter and not payment_pending:
        criteria = _date_criteria(query, today=today)
        metadata.update(criteria)
    else:
        metadata.update({
            "period": "not_applicable" if report == "inventory" else "unfiltered",
            "start_date": None,
            "end_date": None,
            "period_label": "当前库存，不按销售日期筛选" if report == "inventory" else "当前全部可见收款；时间口径待确认Q07",
        })
    metadata["filter_summary"] = _filter_summary(query, metadata)
    return metadata


def _money(value):
    return "" if value is None else f"{Decimal(value):.3f}"


def _local_datetime(value):
    return value.astimezone(REPORT_TIME_ZONE).isoformat(timespec="seconds") if value else ""


def _inventory_data(actor, query):
    workbench = query.get("workbench", "")
    if workbench:
        _reject_workbench_filters(query, ("brand_id", "sku_id", "warehouse_id", "low_stock", "order_id", "scope", "series"))
    items = inventory_rows(actor, query)
    balance_ids = [int(item["id"]) for item in items]
    balances = {
        row.pk: row
        for row in (
            InventoryBalance.objects.filter(pk__in=balance_ids)
            .select_related("sku__product__brand", "sku__product__language")
            .prefetch_related("holds")
        )
    }
    occupied_counts = {
        (row["warehouse_id"], row["sku_id"]): row["order_count"]
        for row in SalesOrderLine.objects.filter(
            active=True,
            order__status__in=(SalesOrder.Status.AWAITING_PAYMENT, SalesOrder.Status.READY_TO_SHIP),
        ).values("warehouse_id", "sku_id").annotate(order_count=Count("order_id", distinct=True))
    }
    available_by_sku = defaultdict(int)
    sku_groups = {}
    combination_labels = {}
    for item in items:
        balance = balances[int(item["id"])]
        sku, product = balance.sku, balance.sku.product
        if workbench and sku.is_active and product.is_active and product.brand.is_active:
            combination = (product.brand_id, product.language_id)
            available_by_sku[sku.pk] += item["available"]
            sku_groups[sku.pk] = combination
            combination_labels[combination] = {
                "brand_id": product.brand_id,
                "brand_name": product.brand.name,
                "language_id": product.language_id,
                "language_name": product.language.name if product.language_id else "未标语言",
            }
        item["barcode"] = balance.sku.barcode
        item["occupied_order_count"] = occupied_counts.get((balance.warehouse_id, balance.sku_id), 0)
        reasons = []
        if item["occupied"]:
            reasons.append(f"订单占用 {item['occupied']} {item['unit']}，涉及{item['occupied_order_count']}个未交付订单")
        for hold in balance.holds.all():
            if hold.remaining_quantity <= 0:
                continue
            label = hold.get_type_display()
            reason = f"（{hold.reason}）" if hold.reason else ""
            reasons.append(f"{label}{reason} {hold.remaining_quantity} {item['unit']}")
        if item["damaged"]:
            reasons.append(f"破损 {item['damaged']} {item['unit']}")
        item["lock_reasons"] = "；".join(reasons) or "无限制"
    columns = [
        ("warehouse", "仓库", "text"), ("brand", "IP", "text"), ("product", "产品", "text"),
        ("code", "SKU编号", "text"), ("barcode", "包装条码", "text"), ("unit", "售卖单位", "text"),
        ("on_hand", "在库数量", "integer"), ("available", "可售数量", "integer"),
        ("occupied", "订单占用", "integer"), ("occupied_order_count", "未交付订单数", "integer"),
        ("reserved", "客户预留", "integer"), ("hoarded", "囤货锁定", "integer"),
        ("damaged", "破损数量", "integer"), ("lock_reasons", "限制原因", "text"),
        ("cost_amount", "库存成本金额", "amount"), ("average_cost", "平均单位成本", "amount"),
    ]
    quantities_by_unit = defaultdict(lambda: {field: 0 for field in ("on_hand", "available", "occupied", "reserved", "damaged")})
    for item in items:
        quantities = quantities_by_unit[item["unit"]]
        for field in quantities:
            quantities[field] += item[field]
    summary = {
        "by_unit": [
            {"unit": unit, **quantities_by_unit[unit]}
            for unit in sorted(quantities_by_unit)
        ],
    }
    if workbench:
        in_stock_combinations = {
            sku_groups[sku_id] for sku_id, available in available_by_sku.items() if available > 0
        }
        summary["workbench"] = {
            "in_stock_ip_language_combination_count": len(in_stock_combinations),
            "in_stock_ip_language_combinations_preview": [
                combination_labels[combination]
                for combination in sorted(
                    in_stock_combinations,
                    key=lambda value: (value[0], value[1] is not None, value[1] or 0),
                )[:3]
            ],
        }
    return items, columns, summary


def _sales_date_query(orders, period):
    if period["start_at"]:
        return orders.filter(
            Q(delivery_mode=SalesOrder.DeliveryMode.SHIPMENT, shipped_at__gte=period["start_at"], shipped_at__lt=period["end_before"])
            | Q(delivery_mode=SalesOrder.DeliveryMode.ON_SITE, completed_at__gte=period["start_at"], completed_at__lt=period["end_before"])
        )
    return orders.filter(
        Q(delivery_mode=SalesOrder.DeliveryMode.SHIPMENT, shipped_at__isnull=False)
        | Q(delivery_mode=SalesOrder.DeliveryMode.ON_SITE, completed_at__isnull=False)
    )


def _sales_data(actor, query, *, today=None, now=None):
    scope = query.get("scope", "delivered") or "delivered"
    if scope not in ("delivered", "unshipped"):
        raise OperationError("销售范围无效。", fields={"scope": "请选择已交付销售或未交付订单。"})
    series_kind = query.get("series", "")
    if series_kind not in ("", "day"):
        raise OperationError("销售序列类型无效。", fields={"series": "请选择day或移除序列筛选。"})
    now = now or _now_local()
    today = today or now.date()
    period = _date_criteria(query, today=today)
    workbench = query.get("workbench", "")
    if workbench:
        if workbench not in {"customer_ranking", "day_over_day"}:
            raise OperationError("工作台统计标记无效。", fields={"workbench": "销售报表不支持该标记。"})
        if query.get("scope") != "delivered":
            raise OperationError("工作台销售统计仅支持已交付销售。", fields={"scope": "请使用delivered。"})
        if query.get("period") != "custom":
            raise OperationError("工作台销售统计需要自定义日期。", fields={"period": "请使用custom。"})
        _reject_workbench_filters(
            query,
            ("sku_id", "brand_id", "warehouse_id", "salesperson_id", "customer_id", "channel_id", "province", "city"),
        )
        if workbench == "customer_ranking":
            if series_kind:
                raise OperationError("客户排行不支持序列参数。", fields={"series": "请移除series。"})
            if period["start_date"] != today.replace(day=1).isoformat() or period["end_date"] != today.isoformat():
                raise OperationError("客户排行仅支持本月累计日期。", fields={"start_date": "请使用北京时间本月1日至今天。", "end_date": "请使用北京时间本月1日至今天。"})
        else:
            expected_start = (today - timedelta(days=1)).isoformat()
            if series_kind != "day":
                raise OperationError("增速统计需要逐日序列。", fields={"series": "请使用day。"})
            if period["start_date"] != expected_start or period["end_date"] != today.isoformat():
                raise OperationError("增速统计仅支持昨天至今天。", fields={"start_date": "请使用北京时间昨天至今天。", "end_date": "请使用北京时间昨天至今天。"})
    series_start = series_end = None
    if series_kind == "day":
        if scope != "delivered" or period["period"] != "custom":
            raise OperationError("逐日序列需要已交付销售的自定义日期范围。", fields={"series": "请使用已交付销售及自定义起止日期。"})
        series_start = date.fromisoformat(period["start_date"])
        series_end = date.fromisoformat(period["end_date"])
        if (series_end - series_start).days + 1 > 366:
            raise OperationError("逐日序列最多查询366天。", fields={"end_date": "请将自定义日期范围缩小到366天以内。"})
    if scope == "unshipped" and period["period"] != "all":
        raise OperationError("未交付订单没有交付日期，不能按交付周期筛选。", fields={"period": "请将周期设为全部；未交付订单按其他条件查询。"})
    orders = _order_filters(_order_queryset(actor), query).exclude(status=SalesOrder.Status.CANCELLED)
    if scope == "unshipped":
        orders = orders.filter(status__in=(SalesOrder.Status.AWAITING_PAYMENT, SalesOrder.Status.READY_TO_SHIP))
    else:
        orders = _sales_date_query(orders, period)
    lines = SalesOrderLine.objects.filter(order_id__in=orders.values("pk"), active=True).select_related(
        "sku__product__brand", "warehouse", "order__salesperson", "order__customer", "order__channel"
    )
    for field, model_field in (("brand_id", "sku__product__brand_id"), ("sku_id", "sku_id"), ("warehouse_id", "warehouse_id")):
        value = _filter_identifier(query, field)
        if value is not None:
            lines = lines.filter(**{model_field: value})
    groups = {}
    daily_revenue = defaultdict(lambda: ZERO) if series_kind == "day" else None
    workbench_daily_revenue = defaultdict(lambda: ZERO) if workbench == "day_over_day" else None
    customer_revenue = {} if workbench == "customer_ranking" else None
    for line in lines.iterator():
        order, sku = line.order, line.sku
        delivered_at = None
        if daily_revenue is not None or customer_revenue is not None:
            delivered_at = order.shipped_at if order.delivery_mode == SalesOrder.DeliveryMode.SHIPMENT else order.completed_at
        if daily_revenue is not None:
            day = delivered_at.astimezone(REPORT_TIME_ZONE).date().isoformat()
            daily_revenue[day] += line.revenue_amount
            if workbench_daily_revenue is not None and delivered_at <= now:
                workbench_daily_revenue[day] += line.revenue_amount
        if customer_revenue is not None and order.customer_id and delivered_at <= now:
            customer = customer_revenue.setdefault(
                order.customer_id,
                {"amount": ZERO, "latest_order": None, "customer_name": ""},
            )
            customer["amount"] += line.revenue_amount
            latest_order = (delivered_at, order.pk)
            if customer["latest_order"] is None or latest_order > customer["latest_order"]:
                customer["latest_order"] = latest_order
                customer["customer_name"] = order.customer_name_snapshot
        key = (
            sku.pk, line.warehouse_id, order.cost_basis, order.salesperson_id, order.customer_id, order.channel_id,
            order.province_snapshot, order.city_snapshot,
        )
        if key not in groups:
            groups[key] = {
                "row_type": "未交付订单" if scope == "unshipped" else "已交付销售",
                "cost_basis": order.cost_basis,
                "cost_basis_label": "出库实际成本" if order.cost_basis == SalesOrder.CostBasis.ACTUAL_OUTFLOW else "历史下单成本快照",
                "sku_id": str(sku.pk), "sku_code": sku.code, "barcode": sku.barcode,
                "brand": sku.product.brand.name, "product": sku.product.name,
                "warehouse_id": str(line.warehouse_id), "warehouse": line.warehouse.name,
                "salesperson_id": str(order.salesperson_id), "salesperson": order.salesperson.get_username(),
                "customer_id": str(order.customer_id) if order.customer_id else "",
                "customer": order.customer_name_snapshot or "门店散客", "channel_id": str(order.channel_id),
                "channel": order.channel.name, "province": order.province_snapshot, "city": order.city_snapshot,
                "unit": sku.unit, "quantity": 0, "order_count": 0,
                "revenue_amount": ZERO, "cost_amount": ZERO, "freight_amount": ZERO, "margin_amount": ZERO,
                "_cost_incomplete": False, "_orders": set(),
            }
        item = groups[key]
        item["quantity"] += line.quantity
        item["_orders"].add(order.pk)
        item["revenue_amount"] += line.revenue_amount
        item["freight_amount"] += line.outbound_freight
        if line.order_cost_amount is None:
            item["_cost_incomplete"] = True
        else:
            item["cost_amount"] += line.order_cost_amount
            item["margin_amount"] += line.revenue_amount - line.order_cost_amount - line.outbound_freight
    items = []
    for item in groups.values():
        item["order_count"] = len(item.pop("_orders"))
        if item.pop("_cost_incomplete"):
            item["cost_amount"] = item["margin_amount"] = None
        for field in ("revenue_amount", "freight_amount"):
            item[field] = _money(item[field])
        for field in ("cost_amount", "margin_amount"):
            item[field] = _money(item[field])
        items.append(item)
    items.sort(key=lambda row: (row["brand"], row["sku_code"], row["warehouse"], row["salesperson"], row["customer"]))
    revenue_label = "订单金额（未交付）" if scope == "unshipped" else "销售金额"
    margin_label = "预估毛利（未交付）" if scope == "unshipped" else "毛利"
    columns = [
        ("sku_code", "SKU编号", "text"), ("barcode", "包装条码", "text"), ("brand", "IP", "text"),
        ("product", "产品", "text"), ("warehouse", "仓库", "text"), ("salesperson", "销售", "text"),
        ("customer", "客户", "text"), ("channel", "渠道", "text"), ("province", "收货省份", "text"),
        ("city", "收货城市", "text"), ("unit", "售卖单位", "text"), ("cost_basis_label", "成本口径", "text"),
        ("quantity", "数量", "integer"), ("order_count", "订单数", "integer"),
        ("revenue_amount", revenue_label, "amount"), ("cost_amount", "成本", "amount"),
        ("freight_amount", "出库运费", "amount"), ("margin_amount", margin_label, "amount"),
    ]
    totals = {}
    for key in ("revenue_amount", "freight_amount"):
        totals[key] = _money(sum((Decimal(row[key]) for row in items), ZERO))
    for key in ("cost_amount", "margin_amount"):
        totals[key] = "" if any(row[key] == "" for row in items) else _money(sum((Decimal(row[key]) for row in items), ZERO))
    quantity_by_unit = defaultdict(int)
    for row in items:
        quantity_by_unit[row["unit"]] += row["quantity"]
    summary = {
        "totals": totals,
        "quantity_by_unit": [{"unit": unit, "quantity": quantity} for unit, quantity in sorted(quantity_by_unit.items())],
        "scope": scope,
    }
    if customer_revenue is not None:
        summary["workbench"] = {
            "customer_ranking": [
                {
                    "customer_id": str(customer_id),
                    "customer_name": customer["customer_name"],
                    "amount": _money(customer["amount"]),
                }
                for customer_id, customer in sorted(
                    customer_revenue.items(), key=lambda entry: (-entry[1]["amount"], entry[0])
                )[:3]
            ],
        }
    elif workbench == "day_over_day":
        today_amount = workbench_daily_revenue[period["end_date"]]
        yesterday_amount = workbench_daily_revenue[period["start_date"]]
        if yesterday_amount:
            with localcontext() as context:
                context.prec = 50
                growth = ((today_amount / yesterday_amount - Decimal("1")) * Decimal("100")).quantize(
                    Decimal("0.01"), rounding=ROUND_HALF_UP
                )
            growth_percent = f"{growth:.2f}"
        else:
            growth_percent = None
        summary["workbench"] = {
            "day_over_day": {
                "today_amount": _money(today_amount),
                "yesterday_amount": _money(yesterday_amount),
                "growth_percent": growth_percent,
            },
        }
    series = None
    if daily_revenue is not None:
        series = []
        for offset in range((series_end - series_start).days + 1):
            day = (series_start + timedelta(days=offset)).isoformat()
            series.append({"date": day, "revenue_amount": _money(daily_revenue[day])})
    return items, columns, summary, series


def _payment_data(actor, query, *, today=None):
    period = _date_criteria(query, today=today)
    orders = _order_filters(_order_queryset(actor), query)
    search_text = query.get("q", "").strip()
    if search_text:
        search = Q(customer_name_snapshot__icontains=search_text)
        if search_text.isascii() and search_text.isdigit() and len(search_text) <= 19:
            order_id = int(search_text)
            if order_id <= 2**63 - 1:
                search |= Q(pk=order_id)
        orders = orders.filter(search)
    records = PaymentRecord.objects.filter(order_id__in=orders.values("pk")).select_related(
        "order__salesperson", "order__customer", "order__channel", "account", "registered_by"
    ).prefetch_related("images")
    if period["start_at"]:
        records = records.filter(
            first_evidence_registered_at__gte=period["start_at"],
            first_evidence_registered_at__lt=period["end_before"],
        )
    status_value = query.get("status", "")
    valid_statuses = dict(PaymentRecord.Status.choices)
    if status_value and status_value not in valid_statuses:
        raise OperationError("核对状态无效。", fields={"status": "请选择待核对、待处理或核对通过。"})
    if status_value:
        records = records.filter(status=status_value)
    account_id = _filter_identifier(query, "account_id")
    if account_id is not None:
        records = records.filter(account_id=account_id)
    paused = query.get("paused", "")
    if paused not in ("", "true", "false"):
        raise OperationError("暂停筛选值无效。", fields={"paused": "请选择全部、暂停或未暂停。"})
    if paused:
        records = records.filter(paused=(paused == "true"))
    items = []
    for payment in records.order_by("order_id").iterator(chunk_size=200):
        order = payment.order
        difference = payment.received_amount - payment.expected_amount
        items.append({
            "order_id": str(order.pk), "submitted_at": _local_datetime(order.created_at),
            "salesperson": order.salesperson.get_username(), "customer": order.customer_name_snapshot or "门店散客",
            "channel": order.channel.name, "order_status": order.status,
            "order_status_label": order.get_status_display(), "status": payment.status,
            "status_label": payment.get_status_display(),
            "paused": "是" if payment.paused else "否", "order_amount": _money(order.revenue_amount),
            "expected_amount": _money(payment.expected_amount), "registered_amount": _money(payment.received_amount),
            "difference_amount": _money(difference), "account_id": str(payment.account_id),
            "account": payment.account.name, "account_method": payment.account.get_method_display(),
            "responsibility": payment.account.get_responsibility_display(),
            "registered_by": payment.registered_by.get_username(),
            "first_evidence_registered_at": _local_datetime(payment.first_evidence_registered_at),
            "actual_received_at": _local_datetime(payment.received_at),
            "image_ids": ", ".join(str(image.pk) for image in payment.images.all()),
            "difference_reason": payment.difference_reason, "problem_reason": payment.problem_reason,
        })
    columns = [
        ("order_id", "订单编号", "text"), ("submitted_at", "订单提交时间", "text"),
        ("salesperson", "销售", "text"), ("customer", "客户快照", "text"), ("channel", "渠道", "text"),
        ("order_status_label", "履约状态", "text"),
        ("status_label", "核对状态", "text"), ("paused", "出库暂停", "text"),
        ("order_amount", "订单金额", "amount"), ("expected_amount", "登记应收金额", "amount"),
        ("registered_amount", "登记收款金额（非实收核实）", "amount"),
        ("difference_amount", "登记差额（收款−应收）", "amount"),
        ("account", "收款账户", "text"), ("account_method", "账户方式", "text"),
        ("responsibility", "核对责任", "text"), ("registered_by", "登记人", "text"),
        ("first_evidence_registered_at", "首次登记收款凭证时间", "text"),
        ("actual_received_at", "登记的实际收款时间（可空）", "text"),
        ("image_ids", "受控凭证图片编号", "text"), ("difference_reason", "金额差异原因", "text"),
        ("problem_reason", "待处理原因", "text"),
    ]
    return items, columns, {}


def _commission_status(order, review):
    payment = getattr(order, "payment_record", None)
    eligible = bool(
        order.status == SalesOrder.Status.COMPLETED
        and payment is not None
        and payment.status == PaymentRecord.Status.PASSED
        and not payment.paused
        and payment.expected_payment_version == order.payment_version
        and payment.verified_version == order.payment_version
    )
    if review is None or review.approved_amount is None:
        return ("pending_review", "待复核") if eligible else ("estimate", "预估")
    if review.approved_amount_version != order.amount_version:
        return "needs_review", "待重审"
    return "approved", "已计入"


def _commission_data(actor, query, *, today=None):
    period = _date_criteria(query, today=today)
    orders = _order_filters(_order_queryset(actor), query).exclude(status=SalesOrder.Status.CANCELLED)
    if not actor.is_superuser and not actor.has_perm("crm.view_all_commissions"):
        orders = orders.filter(salesperson_id=actor.pk)
    commission_status = query.get("commission_status", "")
    valid_states = {"estimate", "pending_review", "approved", "needs_review"}
    if commission_status and commission_status not in valid_states:
        raise OperationError("提成状态无效。", fields={"commission_status": "请选择有效提成状态。"})
    review_time = "commission_review__first_earned_at"
    if period["start_at"]:
        orders = orders.filter(
            Q(**{f"{review_time}__gte": period["start_at"], f"{review_time}__lt": period["end_before"]})
            | Q(**{f"{review_time}__isnull": True, "created_at__gte": period["start_at"], "created_at__lt": period["end_before"]})
        )
    order_items = []
    summaries = defaultdict(lambda: {"count": 0, "amount": ZERO})
    for order in orders.prefetch_related("commission_review", "payment_record").order_by("id").iterator(chunk_size=200):
        review = getattr(order, "commission_review", None)
        state, state_label = _commission_status(order, review)
        if commission_status and state != commission_status:
            continue
        on_date = review.first_earned_at if review and review.first_earned_at else order.created_at
        item = {
            "entry_type": "order", "entry_type_label": "订单提成",
            "effective_at": _local_datetime(on_date), "order_id": str(order.pk),
            "salesperson": order.salesperson.get_username(), "customer": order.customer_name_snapshot or "门店散客",
            "channel": order.channel.name, "status": state, "status_label": state_label,
            "estimate_amount": _money(order.commission_estimate),
            "approved_amount": _money(review.approved_amount) if review and review.approved_amount is not None else "",
            "review_delta": _money(order.commission_estimate - review.approved_amount)
            if review and review.approved_amount is not None and order.commission_estimate is not None and state == "needs_review" else "",
            "adjustment_kind": "", "old_amount": "", "new_amount": "", "delta_amount": "",
            "completion_status": "", "completed_at": "", "completed_by": "",
        }
        order_items.append(item)
        summaries[state]["count"] += 1
        if state in ("estimate", "pending_review") and order.commission_estimate is not None:
            summaries[state]["amount"] += order.commission_estimate
        elif review and review.approved_amount is not None:
            summaries[state]["amount"] += review.approved_amount

    adjustments = CommissionAdjustment.objects.filter(
        order_id__in=_order_filters(_order_queryset(actor), query)
        .exclude(status=SalesOrder.Status.CANCELLED)
        .values("pk"),
        kind__in=(CommissionAdjustment.Kind.REREVIEW, CommissionAdjustment.Kind.COST_CORRECTION),
    ).select_related("order__salesperson", "order__channel", "order__commission_review", "order__payment_record", "completed_by")
    if not actor.is_superuser and not actor.has_perm("crm.view_all_commissions"):
        adjustments = adjustments.filter(order__salesperson_id=actor.pk)
    if period["start_at"]:
        adjustments = adjustments.filter(confirmed_at__gte=period["start_at"], confirmed_at__lt=period["end_before"])
    items = list(order_items)
    adjustment_delta = ZERO
    for adjustment in adjustments.order_by("confirmed_at", "id").iterator():
        state, state_label = _commission_status(adjustment.order, getattr(adjustment.order, "commission_review", None))
        if commission_status and state != commission_status:
            continue
        adjustment_delta += adjustment.delta
        items.append({
            "entry_type": "adjustment", "entry_type_label": "提成调整差额（独立展示）",
            "effective_at": _local_datetime(adjustment.confirmed_at), "order_id": str(adjustment.order_id),
            "salesperson": adjustment.order.salesperson.get_username(), "customer": adjustment.order.customer_name_snapshot or "门店散客",
            "channel": adjustment.order.channel.name, "status": state, "status_label": state_label,
            "estimate_amount": "", "approved_amount": "", "review_delta": "",
            "adjustment_kind": adjustment.get_kind_display(), "old_amount": _money(adjustment.old_amount),
            "new_amount": _money(adjustment.new_amount), "delta_amount": _money(adjustment.delta),
            "completion_status": "已完成" if adjustment.completed_at else "待完成",
            "completed_at": _local_datetime(adjustment.completed_at),
            "completed_by": adjustment.completed_by.get_username() if adjustment.completed_by else "",
        })
    items.sort(key=lambda row: (row["effective_at"], row["entry_type"], row["order_id"]))
    columns = [
        ("effective_at", "统计日期", "text"), ("entry_type_label", "记录类型", "text"),
        ("order_id", "订单编号", "text"), ("salesperson", "销售", "text"),
        ("customer", "客户", "text"), ("channel", "渠道", "text"),
        ("status_label", "提成状态", "text"), ("estimate_amount", "当前预估", "amount"),
        ("approved_amount", "最后批准金额（正式值）", "amount"),
        ("review_delta", "待重审估算差额", "amount"), ("adjustment_kind", "调整类型", "text"),
        ("old_amount", "调整前金额", "amount"), ("new_amount", "调整后金额", "amount"),
        ("delta_amount", "调整差额（不并入正式值）", "amount"),
        ("completion_status", "线下处理", "text"), ("completed_at", "完成时间", "text"),
        ("completed_by", "确认人", "text"),
    ]
    summary = {
        "states": {state: {"count": value["count"], "amount": _money(value["amount"])} for state, value in summaries.items()},
        "adjustment_delta_not_added": _money(adjustment_delta),
    }
    return items, columns, summary


def _filter_options(actor, report):
    if report == "inventory":
        balances = InventoryBalance.objects.select_related("warehouse", "sku__product__brand")
        warehouse_rows = balances.values("warehouse_id", "warehouse__name").distinct().order_by("warehouse__name")
        sku_rows = balances.values("sku_id", "sku__code", "sku__product__name", "sku__unit").distinct().order_by("sku__code")
        brand_rows = balances.values("sku__product__brand_id", "sku__product__brand__name").distinct().order_by("sku__product__brand__name")
        warehouses = [{"id": str(row["warehouse_id"]), "label": row["warehouse__name"]} for row in warehouse_rows]
        skus = [{"id": str(row["sku_id"]), "label": f"{row['sku__code']} · {row['sku__product__name']} · {row['sku__unit']}"} for row in sku_rows]
        brands = [{"id": str(row["sku__product__brand_id"]), "label": row["sku__product__brand__name"]} for row in brand_rows]
        return {"warehouses": warehouses, "skus": skus, "brands": brands}
    orders = _order_queryset(actor)
    if report == "commissions" and not actor.is_superuser and not actor.has_perm("crm.view_all_commissions"):
        orders = orders.filter(salesperson_id=actor.pk)
    order_ids = orders.values("pk")
    warehouses = Warehouse.objects.filter(sales_order_lines__order_id__in=order_ids).distinct().order_by("name")
    skus = SKU.objects.filter(sales_order_lines__order_id__in=order_ids).distinct().select_related("product").order_by("code")
    brands = Brand.objects.filter(products__skus__sales_order_lines__order_id__in=order_ids).distinct().order_by("name")
    users = get_user_model().objects.filter(pk__in=orders.values("salesperson_id")).order_by("username")
    customers = orders.exclude(customer_id__isnull=True).values("customer_id", "customer_name_snapshot").distinct().order_by("customer_name_snapshot")
    channels = orders.values("channel_id", "channel__name").distinct().order_by("channel__name")
    provinces = orders.exclude(province_snapshot="").values_list("province_snapshot", flat=True).distinct().order_by("province_snapshot")
    cities = orders.exclude(city_snapshot="").values_list("city_snapshot", flat=True).distinct().order_by("city_snapshot")
    options = {
        "warehouses": [{"id": str(row.pk), "label": row.name} for row in warehouses],
        "skus": [{"id": str(row.pk), "label": f"{row.code} · {row.product.name} · {row.unit}"} for row in skus],
        "brands": [{"id": str(row.pk), "label": row.name} for row in brands],
        "salespeople": [{"id": str(row.pk), "label": row.get_username()} for row in users],
        "customers": [{"id": str(row["customer_id"]), "label": f"{row['customer_name_snapshot']}（客户{row['customer_id']}）"} for row in customers],
        "channels": [{"id": str(row["channel_id"]), "label": row["channel__name"]} for row in channels],
        "provinces": list(provinces), "cities": list(cities),
    }
    if report == "payments":
        accounts = PaymentAccount.objects.filter(payment_records__order_id__in=order_ids).distinct().order_by("name")
        options["accounts"] = [{"id": str(row.pk), "label": row.name} for row in accounts]
    return options


def build_report(actor, report, query):
    if report not in REPORTS:
        raise OperationError("报表不存在。", status=404, code="not_found")
    workbench = query.get("workbench", "")
    if "workbench" in query and workbench not in WORKBENCH_REPORTS.get(report, set()):
        raise OperationError("工作台统计标记无效。", fields={"workbench": "请移除或使用当前报表支持的标记。"})
    now = _now_local()
    series = None
    if report == "inventory":
        if query.get("start_date") or query.get("end_date") or query.get("period") not in (None, "", "all"):
            raise OperationError("库存报表展示当前实时余额，不支持按销售日期筛选。", fields={"period": "请移除日期筛选。"})
        items, columns, summary = _inventory_data(actor, query)
        metadata = _with_query_metadata(report, query, basis="当前库存余额", date_filter=False)
    elif report == "sales":
        items, columns, summary, series = _sales_data(actor, query, today=now.date(), now=now)
        scope = query.get("scope", "delivered") or "delivered"
        basis = "未交付订单（无交付日期）" if scope == "unshipped" else "快递发货时间／现场交付时间"
        metadata = _with_query_metadata(report, query, basis=basis, today=now.date())
        if scope == "unshipped":
            metadata["date_filter_available"] = False
            metadata["period_label"] = "未交付订单（周期筛选不适用）"
    elif report == "payments":
        items, columns, summary = _payment_data(actor, query, today=now.date())
        metadata = _with_query_metadata(report, query, basis="首次登记收款凭证时间（北京时间）", today=now.date())
    else:
        items, columns, summary = _commission_data(actor, query, today=now.date())
        metadata = _with_query_metadata(report, query, basis="未首次计入按订单提交时间；正式计入按首次计入时间；调整差额按批准时间", today=now.date())
    metadata["queried_at"] = now.isoformat(timespec="seconds")
    metadata["filter_summary"] = _filter_summary(query, metadata)
    result = {
        "name": report,
        "title": REPORTS[report]["title"],
        "items": items,
        "columns": columns,
        "summary": summary,
        "metadata": metadata,
        "filter_options": _filter_options(actor, report),
    }
    if series is not None:
        result["series"] = series
    return result


def page_report_data(query, data, page):
    params = query.copy() if hasattr(query, "copy") else dict(query)
    params["page"] = page
    page_result = paginate(data["items"], params)
    columns = data["columns"]
    result = {
        **data,
        "items": page_result["items"],
        "page": page_result["page"],
        "page_size": page_result["page_size"],
        "total": page_result["total"],
        "has_previous": page_result["has_previous"],
        "has_next": page_result["has_next"],
        "table_rows": [[item.get(key, "") for key, _label, _kind in columns] for item in page_result["items"]],
    }
    return result


def workbook_bytes(data):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    items, columns = data["items"], data["columns"]
    if len(items) + 6 > EXCEL_MAX_ROWS:
        raise OperationError("筛选结果超过单张工作表上限，请缩小筛选范围后导出。", status=413, code="export_too_large")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = data["title"][:31]
    metadata_rows = [
        [data["title"]],
        ["统计范围", data["metadata"]["period_label"]],
        ["统计口径", data["metadata"]["date_basis"]],
        ["筛选条件", data["metadata"]["filter_summary"]],
        ["查询时间（北京时间）", data["metadata"]["queried_at"]],
    ]
    for values in metadata_rows:
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                if len(cell.value) > 32767:
                    raise OperationError("导出说明超过Excel单元格长度上限。", status=422, code="export_text_too_long")
                cell.data_type = "s"
    sheet.append([label for _key, label, _kind in columns])
    for col in sheet[1]:
        col.font = Font(bold=True, size=14)
    for cell in sheet[6]:
        cell.font = Font(bold=True)
    for item in items:
        row = []
        for key, _label, kind in columns:
            value = item.get(key, "")
            if value in (None, ""):
                row.append(None)
            elif kind == "amount":
                try:
                    amount = Decimal(str(value)).quantize(Decimal("0.001"))
                except (InvalidOperation, ValueError):
                    raise OperationError("报表金额无法安全导出。", status=422, code="invalid_export_amount") from None
                significant = amount.normalize().as_tuple().digits if amount else (0,)
                if len(significant) > 15:
                    raise OperationError("有金额超过Excel 15位有效数字精度，未生成可能失真的文件。请缩小筛选范围或先核对金额。", status=422, code="export_amount_precision")
                row.append(amount)
            elif kind == "integer":
                row.append(int(value))
            else:
                text_value = str(value)
                if len(text_value) > 32767:
                    raise OperationError("有文本超过Excel单元格长度上限，未截断导出。请缩小筛选范围或检查该记录。", status=422, code="export_text_too_long")
                row.append(text_value)
        sheet.append(row)
    for row in sheet.iter_rows(min_row=7):
        for cell, (_key, _label, kind) in zip(row, columns):
            if kind == "amount":
                cell.number_format = '#,##0.000;[Red]-#,##0.000'
            elif kind == "text":
                cell.number_format = "@"
                cell.data_type = "s"
    sheet.freeze_panes = "A7"
    sheet.auto_filter.ref = f"A6:{sheet.cell(row=max(6, len(items) + 6), column=len(columns)).coordinate}"
    for column_index, (key, label, _kind) in enumerate(columns, start=1):
        width = min(40, max(12, len(label) + 2, *(min(40, len(str(item.get(key, "")))) for item in items[:100])))
        sheet.column_dimensions[sheet.cell(row=6, column=column_index).column_letter].width = width
    from io import BytesIO

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
