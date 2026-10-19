from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
from uuid import uuid4
from zoneinfo import ZoneInfo

from django.contrib.auth.models import Permission
from django.test import SimpleTestCase
from django.utils import timezone
from openpyxl import load_workbook

from .models import (
    Brand,
    BusinessControl,
    BusinessOperation,
    CommissionAdjustment,
    CommissionReview,
    InventoryBalance,
    InventoryHold,
    PaymentRecord,
    Product,
    ProductLanguageOption,
    ReceiptImage,
    SalesOrder,
    SKU,
    Warehouse,
)
from .operations import OperationError
from .reports import _date_criteria, build_report, workbook_bytes
from .test_phase4 import Phase4Base
from .fulfillment import deliver_on_site, register_payment, ship_sales_order


SHANGHAI = ZoneInfo("Asia/Shanghai")


class ReportDateCriteriaTests(SimpleTestCase):
    def test_rolling_month_clamps_month_end_and_keeps_beijing_bounds(self):
        result = _date_criteria({"period": "6m"}, today=date(2026, 8, 31))
        self.assertEqual(result["start_date"], "2026-02-28")
        self.assertEqual(result["end_date"], "2026-08-31")
        self.assertEqual(result["start_at"], datetime(2026, 2, 28, tzinfo=SHANGHAI))
        self.assertEqual(result["end_before"], datetime(2026, 9, 1, tzinfo=SHANGHAI))

    def test_leap_day_and_custom_range_validation(self):
        leap = _date_criteria({"period": "6m"}, today=date(2024, 8, 31))
        self.assertEqual(leap["start_date"], "2024-02-29")
        custom = _date_criteria({"period": "custom", "start_date": "2026-02-28", "end_date": "2026-03-01"})
        self.assertEqual(custom["start_date"], "2026-02-28")
        with self.assertRaises(OperationError):
            _date_criteria({"period": "custom", "start_date": "2026-03-02", "end_date": "2026-03-01"})
        with self.assertRaises(OperationError):
            _date_criteria({"period": "custom", "start_date": "2026-02-30", "end_date": "2026-03-01"})
        with self.assertRaises(OperationError) as overflow:
            _date_criteria({"period": "custom", "start_date": "9999-12-31", "end_date": "9999-12-31"})
        self.assertIn("end_date", overflow.exception.fields)


class ReportPhase6Tests(Phase4Base):
    def setUp(self):
        super().setUp()
        self.balance.on_hand = 20
        self.balance.cost_amount = Decimal("20.000")
        self.balance.save(update_fields=("on_hand", "cost_amount"))

    def _set_sales_date(self, order, field, value):
        SalesOrder.objects.filter(pk=order.pk).update(**{field: value})

    def _create_report_order(self, amount, delivered_at, *, customer=None, customer_name=None, actor=None, delivery_mode=SalesOrder.DeliveryMode.ON_SITE):
        address_snapshot = {
            "recipient": "Synthetic Recipient",
            "contact": "Synthetic Contact",
            "province": "北京市",
            "city": "北京市",
            "detail": "Synthetic street 1",
        } if delivery_mode == SalesOrder.DeliveryMode.SHIPMENT else {}
        payload = self.order_payload(
            delivery_mode=delivery_mode,
            address_snapshot=address_snapshot,
            lines=[{"sku_id": self.sku.pk, "warehouse_id": self.warehouse.pk, "quantity": 1, "sale_price": str(amount)}],
        )
        if customer:
            payload["customer_id"] = customer.pk
        _, _, order = self.preview_and_confirm(payload, actor=actor)
        fields = {"status": SalesOrder.Status.COMPLETED}
        fields["shipped_at" if delivery_mode == SalesOrder.DeliveryMode.SHIPMENT else "completed_at"] = delivered_at
        if customer_name is not None:
            fields["customer_name_snapshot"] = customer_name
        SalesOrder.objects.filter(pk=order.pk).update(**fields)
        return order

    def test_workbench_customer_ranking_is_opt_in_scoped_and_snapshot_based(self):
        customer_a = self.create_customer(name="Ranking A")
        customer_b = self.create_customer(name="Ranking B")
        customer_c = self.create_customer(name="Ranking C")
        customer_d = self.create_customer(name="Ranking D")
        other_customer = self.create_customer(self.sales_two, name="Other Sales Customer")
        self._create_report_order("40.000", datetime(2026, 10, 2, 12, tzinfo=SHANGHAI), customer=customer_a, customer_name="A旧快照")
        latest_a = self._create_report_order(
            "60.000", datetime(2026, 10, 8, 12, tzinfo=SHANGHAI), customer=customer_a,
            customer_name="A新快照", delivery_mode=SalesOrder.DeliveryMode.SHIPMENT,
        )
        self._create_report_order("90.000", datetime(2026, 10, 7, 12, tzinfo=SHANGHAI), customer=customer_b, customer_name="B快照")
        self._create_report_order("80.000", datetime(2026, 10, 6, 12, tzinfo=SHANGHAI), customer=customer_c, customer_name="C快照")
        self._create_report_order("70.000", datetime(2026, 10, 5, 12, tzinfo=SHANGHAI), customer=customer_d, customer_name="D快照")
        self._create_report_order("999.000", datetime(2026, 10, 8, 12, tzinfo=SHANGHAI), customer_name="门店散客")
        self._create_report_order("500.000", datetime(2026, 10, 8, 12, tzinfo=SHANGHAI), customer=other_customer, actor=self.sales_two)
        self._create_report_order("5000.000", datetime(2026, 10, 9, 20, tzinfo=SHANGHAI), customer=customer_a, customer_name="未来快照")
        self._create_report_order("2000.000", datetime(2026, 9, 30, 12, tzinfo=SHANGHAI), customer=customer_a)

        query = {
            "period": "custom", "start_date": "2026-10-01", "end_date": "2026-10-09",
            "scope": "delivered", "page_size": 1,
        }
        now = datetime(2026, 10, 9, 12, tzinfo=SHANGHAI)
        with patch("crm.reports._now_local", return_value=now):
            self.assertEqual(self.client.get("/api/v1/reports/sales", query).status_code, 401)
            no_permission = self.User.objects.create_user(username="report-no-permission", password="test-only-password")
            self.client.force_login(no_permission)
            self.assertEqual(self.client.get("/api/v1/reports/sales", query).status_code, 403)

            self.client.force_login(self.sales)
            default = self.client.get("/api/v1/reports/sales", query)
            opted_in = self.client.get("/api/v1/reports/sales", {**query, "workbench": "customer_ranking"})
            self.assertEqual((default.status_code, opted_in.status_code), (200, 200))
            default_data, data = default.json(), opted_in.json()
            self.assertNotIn("workbench", default_data["summary"])
            self.assertEqual(data["summary"].get("totals"), default_data["summary"].get("totals"))
            self.assertEqual(data["items"], default_data["items"])
            self.assertEqual(data["columns"], default_data["columns"])
            self.assertEqual(data["total"], default_data["total"])
            self.assertEqual(data["summary"]["workbench"]["customer_ranking"], [
                {"customer_id": str(customer_a.pk), "customer_name": "A新快照", "amount": "100.000"},
                {"customer_id": str(customer_b.pk), "customer_name": "B快照", "amount": "90.000"},
                {"customer_id": str(customer_c.pk), "customer_name": "C快照", "amount": "80.000"},
            ])

            self.client.force_login(self.owner)
            owner_data = self.client.get("/api/v1/reports/sales", {**query, "workbench": "customer_ranking"}).json()
            self.assertEqual(owner_data["summary"]["workbench"]["customer_ranking"][0]["customer_id"], str(other_customer.pk))

            report_only = self.User.objects.create_user(username="report-without-customer-view", password="test-only-password")
            report_only.user_permissions.add(Permission.objects.get(content_type__app_label="crm", codename="view_salesorder"))
            self.assertFalse(report_only.has_perm("crm.view_customer"))
            SalesOrder.objects.filter(pk=latest_a.pk).update(salesperson=report_only)
            self.client.force_login(report_only)
            report_only_data = self.client.get("/api/v1/reports/sales", {**query, "workbench": "customer_ranking"}).json()
            self.assertEqual(report_only_data["summary"]["workbench"]["customer_ranking"], [
                {"customer_id": str(customer_a.pk), "customer_name": "A新快照", "amount": "60.000"},
            ])

            self.client.force_login(self.sales)
            for invalid in (
                {**query, "workbench": "unknown"},
                {**query, "workbench": ""},
                {**query, "workbench": "ip_language_stock_count"},
                {**query, "workbench": "customer_ranking", "end_date": "2026-10-08"},
                {**query, "workbench": "customer_ranking", "scope": "unshipped"},
                {**query, "workbench": "customer_ranking", "brand_id": str(self.brand.pk)},
            ):
                response = self.client.get("/api/v1/reports/sales", invalid)
                self.assertEqual(response.status_code, 400, response.content)
                self.assertTrue(response.json()["error"]["fields"])

            default_report = build_report(self.sales, "sales", query)
            workbench_report = build_report(self.sales, "sales", {**query, "workbench": "customer_ranking"})
            exported_rows = []
            for report in (default_report, workbench_report):
                workbook = load_workbook(BytesIO(workbook_bytes(report)), read_only=True)
                exported_rows.append(list(workbook.active.iter_rows(min_row=6, values_only=True)))
                workbook.close()
            self.assertEqual(exported_rows[0], exported_rows[1])

    def test_workbench_day_over_day_uses_beijing_dates_and_decimal_rounding(self):
        self._create_report_order(
            "80.000", datetime(2026, 10, 8, 15, 59, 59, 999999, tzinfo=ZoneInfo("UTC")),
            delivery_mode=SalesOrder.DeliveryMode.SHIPMENT,
        )
        self._create_report_order(
            "89.876", datetime(2026, 10, 8, 16, 0, tzinfo=ZoneInfo("UTC")),
            delivery_mode=SalesOrder.DeliveryMode.ON_SITE,
        )
        self._create_report_order("50.000", datetime(2026, 10, 9, 20, tzinfo=SHANGHAI))
        query = {
            "workbench": "day_over_day", "period": "custom", "start_date": "2026-10-08",
            "end_date": "2026-10-09", "scope": "delivered", "series": "day",
        }
        self.client.force_login(self.sales)
        with patch("crm.reports._now_local", return_value=datetime(2026, 10, 9, 12, tzinfo=SHANGHAI)):
            response = self.client.get("/api/v1/reports/sales", query)
            self.assertEqual(response.status_code, 200, response.content)
            data = response.json()
            self.assertEqual(data["series"], [
                {"date": "2026-10-08", "revenue_amount": "80.000"},
                {"date": "2026-10-09", "revenue_amount": "139.876"},
            ])
            self.assertEqual(data["summary"]["workbench"]["day_over_day"], {
                "today_amount": "89.876", "yesterday_amount": "80.000", "growth_percent": "12.35",
            })
            for invalid in (
                {**query, "start_date": "2026-10-07"},
                {**query, "scope": "unshipped"},
                {**query, "series": ""},
            ):
                failed = self.client.get("/api/v1/reports/sales", invalid)
                self.assertEqual(failed.status_code, 400, failed.content)
                self.assertTrue(failed.json()["error"]["fields"])

    def test_workbench_day_over_day_returns_null_growth_for_zero_yesterday(self):
        self._create_report_order("10.000", datetime(2026, 10, 9, 10, tzinfo=SHANGHAI))
        query = {
            "workbench": "day_over_day", "period": "custom", "start_date": "2026-10-08",
            "end_date": "2026-10-09", "scope": "delivered", "series": "day",
        }
        self.client.force_login(self.sales)
        with patch("crm.reports._now_local", return_value=datetime(2026, 10, 9, 12, tzinfo=SHANGHAI)):
            response = self.client.get("/api/v1/reports/sales", query)
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["summary"]["workbench"]["day_over_day"], {
            "today_amount": "10.000", "yesterday_amount": "0.000", "growth_percent": None,
        })

    def test_inventory_workbench_counts_enabled_positive_available_ip_language_pairs(self):
        InventoryBalance.objects.create(sku=self.sku, warehouse=self.second_warehouse, on_hand=2)
        same_language_sku = SKU.objects.create(product=self.product, code="TST-0002", unit="盒")
        InventoryBalance.objects.create(sku=same_language_sku, warehouse=self.second_warehouse, on_hand=3)

        english = ProductLanguageOption.objects.get_or_create(name="英文")[0]
        japanese = ProductLanguageOption.objects.get_or_create(name="日文")[0]
        korean = ProductLanguageOption.objects.get_or_create(name="韩文")[0]
        english_product = Product.objects.create(
            brand=self.brand, type=Product.Type.BOOSTER, name="Active English", language=english,
        )
        english_sku = SKU.objects.create(product=english_product, code="TST-0003", unit="盒")
        InventoryBalance.objects.create(sku=english_sku, warehouse=self.warehouse, on_hand=1)

        for number in (4, 5):
            merchandise = Product.objects.create(
                brand=self.brand, type=Product.Type.MERCHANDISE, name=f"Blank Language {number}",
            )
            sku = SKU.objects.create(product=merchandise, code=f"TST-{number:04d}", unit="个")
            InventoryBalance.objects.create(sku=sku, warehouse=self.warehouse, on_hand=1)

        other_brand = Brand.objects.create(name="Other Active IP", sku_prefix="OTH-")
        other_blank = Product.objects.create(brand=other_brand, type=Product.Type.MERCHANDISE, name="Other Blank")
        other_blank_sku = SKU.objects.create(product=other_blank, code="OTH-0001", unit="个")
        InventoryBalance.objects.create(sku=other_blank_sku, warehouse=self.warehouse, on_hand=1)
        locked_product = Product.objects.create(brand=other_brand, type=Product.Type.BOOSTER, name="Locked Stock", language=japanese)
        locked_sku = SKU.objects.create(product=locked_product, code="OTH-0002", unit="盒")
        InventoryBalance.objects.create(sku=locked_sku, warehouse=self.warehouse, on_hand=2, occupied=2)

        inactive_brand = Brand.objects.create(name="Disabled IP", sku_prefix="OFF-", is_active=False)
        inactive_brand_product = Product.objects.create(
            brand=inactive_brand, type=Product.Type.BOOSTER, name="Disabled Brand Product", language=korean,
        )
        inactive_brand_sku = SKU.objects.create(product=inactive_brand_product, code="OFF-0001", unit="盒")
        InventoryBalance.objects.create(sku=inactive_brand_sku, warehouse=self.warehouse, on_hand=2)
        inactive_product = Product.objects.create(
            brand=self.brand, type=Product.Type.BOOSTER, name="Disabled Product", language=korean, is_active=False,
        )
        inactive_product_sku = SKU.objects.create(product=inactive_product, code="TST-0006", unit="盒")
        InventoryBalance.objects.create(sku=inactive_product_sku, warehouse=self.warehouse, on_hand=2)
        inactive_sku_product = Product.objects.create(
            brand=self.brand, type=Product.Type.BOOSTER, name="Disabled SKU Product", language=japanese,
        )
        inactive_sku = SKU.objects.create(product=inactive_sku_product, code="TST-0007", unit="盒", is_active=False)
        InventoryBalance.objects.create(sku=inactive_sku, warehouse=self.warehouse, on_hand=2)

        query = {"workbench": "ip_language_stock_count", "page_size": 1}
        self.client.force_login(self.sales)
        opted_in = self.client.get("/api/v1/reports/inventory", query)
        default = self.client.get("/api/v1/reports/inventory", {"page_size": 1})
        self.assertEqual((opted_in.status_code, default.status_code), (200, 200))
        opted_in_data, default_data = opted_in.json(), default.json()
        workbench_summary = opted_in_data["summary"]["workbench"]
        self.assertEqual(workbench_summary["in_stock_ip_language_combination_count"], 4)
        expected_preview = [
            {"brand_id": self.brand.pk, "brand_name": self.brand.name, "language_id": self.product.language_id,
             "language_name": self.product.language.name},
            {"brand_id": self.brand.pk, "brand_name": self.brand.name, "language_id": english.pk,
             "language_name": english.name},
            {"brand_id": self.brand.pk, "brand_name": self.brand.name, "language_id": None,
             "language_name": "未标语言"},
            {"brand_id": other_brand.pk, "brand_name": other_brand.name, "language_id": None,
             "language_name": "未标语言"},
        ]
        expected_preview.sort(
            key=lambda row: (row["brand_id"], row["language_id"] is not None, row["language_id"] or 0)
        )
        preview = workbench_summary["in_stock_ip_language_combinations_preview"]
        self.assertEqual(preview, expected_preview[:3])
        self.assertEqual(len(preview), 3)
        self.assertTrue(all(set(row) == {"brand_id", "brand_name", "language_id", "language_name"} for row in preview))
        self.assertNotIn("workbench", default_data["summary"])
        self.assertEqual(opted_in_data["items"], default_data["items"])
        self.assertEqual(opted_in_data["columns"], default_data["columns"])
        self.assertEqual(opted_in_data["total"], default_data["total"])
        self.assertGreater(opted_in_data["total"], 1)

        exported_rows = []
        for report in (
            build_report(self.sales, "inventory", {"page_size": 1}),
            build_report(self.sales, "inventory", query),
        ):
            workbook = load_workbook(BytesIO(workbook_bytes(report)), read_only=True)
            exported_rows.append(list(workbook.active.iter_rows(min_row=6, values_only=True)))
            workbook.close()
        self.assertEqual(exported_rows[0], exported_rows[1])

        no_permission = self.User.objects.create_user(username="inventory-report-no-permission", password="test-only-password")
        self.client.force_login(no_permission)
        self.assertEqual(self.client.get("/api/v1/reports/inventory", query).status_code, 403)
        for invalid in (
            {**query, "workbench": "unknown"},
            {**query, "workbench": ""},
            {**query, "workbench": "customer_ranking"},
            {**query, "workbench": "ip_language_stock_count", "warehouse_id": str(self.warehouse.pk)},
        ):
            self.client.force_login(self.sales)
            response = self.client.get("/api/v1/reports/inventory", invalid)
            self.assertEqual(response.status_code, 400, response.content)
            self.assertTrue(response.json()["error"]["fields"])

    def test_sales_groups_order_lines_per_unit_warehouse_and_snapshot_region(self):
        second_balance = InventoryBalance.objects.create(
            sku=self.sku, warehouse=self.second_warehouse, on_hand=10, cost_amount=Decimal("20.000")
        )
        pack_sku = self.new_sku(2)
        pack_sku.unit = "包"
        pack_sku.save(update_fields=("unit",))
        InventoryBalance.objects.create(sku=pack_sku, warehouse=self.warehouse, on_hand=10, cost_amount=Decimal("20.000"))
        order = self.new_order(delivery_mode=SalesOrder.DeliveryMode.SHIPMENT, lines=[
            {"sku_id": self.sku.pk, "warehouse_id": self.warehouse.pk, "quantity": 1, "sale_price": "5.000"},
            {"sku_id": self.sku.pk, "warehouse_id": self.warehouse.pk, "quantity": 1, "sale_price": "5.000"},
            {"sku_id": self.sku.pk, "warehouse_id": self.second_warehouse.pk, "quantity": 1, "sale_price": "5.000"},
            {"sku_id": pack_sku.pk, "warehouse_id": self.warehouse.pk, "quantity": 2, "sale_price": "5.000"},
        ])
        self.register(order)
        self.verify(order)
        order.refresh_from_db()
        ship_sales_order(self.user, order.pk, self.shipment_data(order))
        shanghai_delivery = datetime(2025, 10, 1, 15, 59, 59, 999999, tzinfo=ZoneInfo("UTC"))
        self._set_sales_date(order, "shipped_at", shanghai_delivery)
        self._set_sales_date(order, "province_snapshot", "浙江省")
        self._set_sales_date(order, "city_snapshot", "杭州市")

        data = build_report(self.sales, "sales", {
            "period": "custom", "start_date": "2025-10-01", "end_date": "2025-10-01",
        })
        self.assertEqual(len(data["items"]), 3)
        first = next(row for row in data["items"] if row["warehouse_id"] == str(self.warehouse.pk))
        second = next(row for row in data["items"] if row["warehouse_id"] == str(self.second_warehouse.pk))
        self.assertEqual(first["quantity"], 2)
        self.assertEqual(first["order_count"], 1)
        self.assertEqual(first["revenue_amount"], "10.000")
        self.assertEqual(first["cost_amount"], "2.000")
        self.assertEqual(first["margin_amount"], "8.000")
        self.assertEqual(second["quantity"], 1)
        self.assertEqual(second["cost_amount"], "2.000")
        self.assertEqual(second["province"], "浙江省")
        self.assertEqual(second["city"], "杭州市")
        self.assertEqual(data["summary"]["quantity_by_unit"], [{"unit": "包", "quantity": 2}, {"unit": "盒", "quantity": 3}])
        filtered = build_report(self.sales, "sales", {
            "period": "custom", "start_date": "2025-10-01", "end_date": "2025-10-01",
            "warehouse_id": str(self.warehouse.pk),
        })
        self.assertEqual(len(filtered["items"]), 2)
        outside = build_report(self.sales, "sales", {
            "period": "custom", "start_date": "2025-10-02", "end_date": "2025-10-02",
        })
        self.assertEqual(outside["items"], [])

        daily_query = {
            "period": "custom",
            "start_date": "2025-09-30",
            "end_date": "2025-10-02",
            "series": "day",
        }
        daily = build_report(self.sales, "sales", daily_query)
        self.assertEqual(daily["series"], [
            {"date": "2025-09-30", "revenue_amount": "0.000"},
            {"date": "2025-10-01", "revenue_amount": "25.000"},
            {"date": "2025-10-02", "revenue_amount": "0.000"},
        ])
        self.assertEqual(
            sum((Decimal(point["revenue_amount"]) for point in daily["series"]), Decimal("0.000")),
            Decimal(daily["summary"]["totals"]["revenue_amount"]),
        )
        self.client.force_login(self.sales)
        series_response = self.client.get("/api/v1/reports/sales", daily_query)
        self.assertEqual(series_response.status_code, 200, series_response.content)
        self.assertEqual(series_response.json()["series"], daily["series"])
        leap_year_series = self.client.get("/api/v1/reports/sales", {
            "period": "custom",
            "start_date": "2024-01-01",
            "end_date": "2024-12-31",
            "series": "day",
        })
        self.assertEqual(leap_year_series.status_code, 200, leap_year_series.content)
        self.assertEqual(len(leap_year_series.json()["series"]), 366)
        overflow_date = self.client.get("/api/v1/reports/sales", {
            "period": "custom",
            "start_date": "9999-12-31",
            "end_date": "9999-12-31",
        })
        self.assertEqual(overflow_date.status_code, 400, overflow_date.content)
        self.assertIn("end_date", overflow_date.json()["error"]["fields"])
        too_many_days = self.client.get("/api/v1/reports/sales", {
            "period": "custom",
            "start_date": "2024-01-01",
            "end_date": "2025-01-01",
            "series": "day",
        })
        self.assertEqual(too_many_days.status_code, 400)
        self.assertIn("end_date", too_many_days.json()["error"]["fields"])
        page = self.client.get("/reports/sales/?period=custom&start_date=2025-10-01&end_date=2025-10-01")
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "杭州市")
        self.assertContains(page, "毛利")
        second_balance.refresh_from_db()
        self.assertEqual(second_balance.on_hand, 9)

    def test_legacy_snapshot_cost_and_preoutflow_blank_sales_export(self):
        legacy = self._create_report_order(
            "8.000", datetime(2025, 10, 1, 12, 0, tzinfo=ZoneInfo("UTC")),
        )
        legacy.cost_basis = SalesOrder.CostBasis.ORDER_SNAPSHOT
        legacy.cost_amount = Decimal("2.000")
        legacy.margin_amount = Decimal("6.000")
        legacy.commission_estimate = Decimal("0.600")
        legacy.save(update_fields=("cost_basis", "cost_amount", "margin_amount", "commission_estimate"))
        legacy_line = legacy.lines.get(active=True)
        legacy_line.unit_cost_snapshot = Decimal("2.000")
        legacy_line.order_cost_amount = Decimal("2.000")
        legacy_line.save(update_fields=("unit_cost_snapshot", "order_cost_amount"))

        legacy_report = build_report(self.sales, "sales", {
            "period": "custom", "start_date": "2025-10-01", "end_date": "2025-10-01",
        })
        legacy_item = next(row for row in legacy_report["items"] if row["sku_id"] == str(self.sku.pk))
        self.assertEqual(legacy_item["cost_basis"], SalesOrder.CostBasis.ORDER_SNAPSHOT)
        self.assertEqual(legacy_item["cost_amount"], "2.000")
        self.assertEqual(legacy_item["margin_amount"], "6.000")

        payload = self.order_payload(lines=[{
            "sku_id": self.sku.pk, "warehouse_id": self.warehouse.pk,
            "quantity": 1, "sale_price": "8.000",
        }])
        _, _, pending = self.preview_and_confirm(payload)
        self.assertEqual(pending.cost_basis, SalesOrder.CostBasis.ACTUAL_OUTFLOW)
        query = {"scope": "unshipped", "period": "all"}
        report = build_report(self.sales, "sales", query)
        item = next(row for row in report["items"] if row["sku_id"] == str(self.sku.pk))
        self.assertEqual(item["cost_basis_label"], "出库实际成本")
        self.assertEqual(item["cost_amount"], "")
        self.assertEqual(item["margin_amount"], "")
        self.assertEqual(report["summary"]["totals"]["cost_amount"], "")
        self.assertEqual(report["summary"]["totals"]["margin_amount"], "")

        self.client.force_login(self.sales)
        api = self.client.get("/api/v1/reports/sales", query)
        self.assertEqual(api.status_code, 200, api.content)
        api_item = next(row for row in api.json()["items"] if row["sku_id"] == str(self.sku.pk))
        self.assertEqual(api_item["cost_amount"], "")
        self.assertEqual(api_item["margin_amount"], "")
        export = self.client.get("/api/v1/reports/sales/export", query)
        self.assertEqual(export.status_code, 200, export.content)
        from io import BytesIO
        from openpyxl import load_workbook

        sheet = load_workbook(BytesIO(export.content), data_only=False).active
        columns = {cell.value: cell.column for cell in sheet[6]}
        self.assertIsNone(sheet.cell(row=7, column=columns["成本"]).value)
        self.assertIsNone(sheet.cell(row=7, column=columns["预估毛利（未交付）"]).value)

    def test_unshipped_orders_are_explicit_and_never_mixed_into_sales(self):
        self.new_order()
        delivered = build_report(self.sales, "sales", {"period": "all"})
        pending = build_report(self.sales, "sales", {"scope": "unshipped", "period": "all"})
        self.assertEqual(delivered["items"], [])
        self.assertEqual(pending["summary"]["scope"], "unshipped")
        self.assertEqual(pending["items"][0]["row_type"], "未交付订单")
        with self.assertRaises(OperationError):
            build_report(self.sales, "sales", {"scope": "unshipped", "period": "1m"})

    def test_inventory_report_keeps_live_balance_and_hoard_reason(self):
        self.balance.hoarded = 2
        self.balance.save(update_fields=("hoarded",))
        InventoryHold.objects.create(
            balance=self.balance, type=InventoryHold.Type.HOARD, initial_quantity=2, remaining_quantity=2,
            reason="Synthetic event stock", created_by=self.user,
        )
        second_balance = InventoryBalance.objects.create(
            sku=self.sku, warehouse=self.second_warehouse, on_hand=5, cost_amount=Decimal("5.000")
        )
        pack_sku = self.new_sku(2)
        pack_sku.unit = "包"
        pack_sku.save(update_fields=("unit",))
        InventoryBalance.objects.create(sku=pack_sku, warehouse=self.warehouse, on_hand=4, cost_amount=Decimal("4.000"))
        data = build_report(self.sales, "inventory", {})
        row = next(item for item in data["items"] if item["id"] == str(self.balance.pk))
        self.assertEqual(row["available"], 18)
        self.assertIn("囤货（Synthetic event stock） 2 盒", row["lock_reasons"])
        self.assertEqual(data["summary"]["by_unit"], [
            {"unit": "包", "on_hand": 4, "available": 4, "occupied": 0, "reserved": 0, "damaged": 0},
            {"unit": "盒", "on_hand": 25, "available": 23, "occupied": 0, "reserved": 0, "damaged": 0},
        ])

        control = BusinessControl.objects.get(pk=1)
        control.low_stock_enabled = True
        control.default_threshold = 5
        control.save(update_fields=("low_stock_enabled", "default_threshold"))
        low_stock = build_report(self.sales, "inventory", {"low_stock": "true"})
        self.assertEqual(low_stock["summary"]["by_unit"], [
            {"unit": "包", "on_hand": 4, "available": 4, "occupied": 0, "reserved": 0, "damaged": 0},
            {"unit": "盒", "on_hand": 5, "available": 5, "occupied": 0, "reserved": 0, "damaged": 0},
        ])
        pack_sku.low_stock_enabled = False
        pack_sku.save(update_fields=("low_stock_enabled",))
        sku_disabled = build_report(self.sales, "inventory", {"low_stock": "true"})
        self.assertEqual(sku_disabled["summary"]["by_unit"], [
            {"unit": "盒", "on_hand": 5, "available": 5, "occupied": 0, "reserved": 0, "damaged": 0},
        ])
        control.low_stock_enabled = False
        control.save(update_fields=("low_stock_enabled",))
        global_disabled = build_report(self.sales, "inventory", {"low_stock": "true"})
        self.assertEqual(global_disabled["items"], [])
        self.assertEqual(global_disabled["summary"]["by_unit"], [])
        self.assertFalse(data["metadata"]["date_filter_available"])
        self.client.force_login(self.sales)
        api = self.client.get("/api/v1/reports/inventory?page_size=1")
        self.assertEqual(api.status_code, 200, api.content)
        self.assertEqual(api.json()["total"], 3)
        self.assertEqual(
            set(api.json()["filter_options"]),
            {"warehouses", "skus", "brands"},
        )
        self.assertEqual(
            {item["id"] for item in api.json()["filter_options"]["warehouses"]},
            {str(self.warehouse.pk), str(self.second_warehouse.pk)},
        )
        self.assertEqual(
            {item["id"] for item in api.json()["filter_options"]["skus"]},
            {str(self.sku.pk), str(pack_sku.pk)},
        )
        page = self.client.get("/reports/inventory/")
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "Synthetic event stock")

    def test_payment_first_evidence_time_survives_failure_replay_updates_and_rechecks(self):
        order = self.new_order()
        image = self.add_payment_image()
        request = {
            "operation_key": str(uuid4()),
            "expected_version": order.version,
            "account_id": self.public_account.pk,
            "received_amount": str(order.revenue_amount),
            "image_ids": [image.pk],
        }
        before = timezone.now()
        with patch("crm.fulfillment._create_operation", side_effect=RuntimeError("synthetic rollback")):
            with self.assertRaises(RuntimeError):
                register_payment(self.sales, order.pk, request)
        self.assertFalse(PaymentRecord.objects.filter(order=order).exists())
        self.assertIsNone(ReceiptImage.objects.get(pk=image.pk).payment_id)
        self.assertFalse(BusinessOperation.objects.filter(sales_order=order, action="payment_register").exists())
        order.refresh_from_db()
        self.assertEqual(order.status, SalesOrder.Status.AWAITING_PAYMENT)

        result = register_payment(self.sales, order.pk, request)
        self.assertEqual(register_payment(self.sales, order.pk, request), result)
        payment = PaymentRecord.objects.get(order=order)
        first_registered_at = payment.first_evidence_registered_at
        operation = BusinessOperation.objects.get(operation_key=request["operation_key"])
        self.assertEqual(first_registered_at, operation.created_at)
        self.assertGreaterEqual(first_registered_at, before)
        self.assertLessEqual(first_registered_at, timezone.now())

        self.verify(order, actor=self.owner)
        second_image = self.add_payment_image()
        self.register(
            order,
            account=self.private_account,
            received_amount="9.000",
            difference_reason="Synthetic correction",
            image=second_image,
        )
        payment.refresh_from_db()
        self.assertEqual(payment.first_evidence_registered_at, first_registered_at)
        self.verify(order, actor=self.owner)
        self.verify(order, actor=self.owner)
        payment.refresh_from_db()
        self.assertEqual(payment.first_evidence_registered_at, first_registered_at)

    def test_payment_report_uses_first_evidence_registration_time(self):
        order = self.new_order()
        self.register(order)
        data = build_report(self.sales, "payments", {})
        self.assertEqual(len(data["items"]), 1)
        self.assertTrue(data["metadata"]["date_filter_available"])
        self.assertFalse(data["metadata"]["payment_date_basis_pending"])
        self.assertEqual(data["metadata"]["date_basis"], "首次登记收款凭证时间（北京时间）")
        self.assertEqual(data["summary"], {})
        self.assertTrue(data["items"][0]["first_evidence_registered_at"])
        self.assertEqual(data["items"][0]["actual_received_at"], "")
        self.assertIn("非实收核实", dict((key, label) for key, label, _ in data["columns"]).get("registered_amount", ""))
        self.client.force_login(self.sales)
        response = self.client.get("/api/v1/reports/payments?period=all")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["date_filter_available"])
        self.assertFalse(response.json()["payment_date_basis_pending"])
        page = self.client.get("/reports/payments/")
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "首次成功登记收款凭证时间")
        self.assertContains(page, "开始日期")
        invalid_page = self.client.get("/reports/payments/?page_size=not-a-number")
        self.assertEqual(invalid_page.status_code, 200)
        self.assertContains(invalid_page, "page与page_size必须为整数")

    def test_payment_report_applies_beijing_boundaries_and_matches_api_page_and_xlsx(self):
        registration_times = [
            datetime(2026, 9, 29, 23, 59, 59, 999999, tzinfo=SHANGHAI),
            datetime(2026, 9, 30, 0, 0, tzinfo=SHANGHAI),
            datetime(2026, 10, 1, 0, 0, tzinfo=SHANGHAI),
            datetime(2026, 10, 1, 23, 59, 59, 999999, tzinfo=SHANGHAI),
            datetime(2026, 10, 2, 0, 0, tzinfo=SHANGHAI),
        ]
        orders = []
        for registered_at in registration_times:
            order = self.new_order()
            self.register(order)
            PaymentRecord.objects.filter(order=order).update(first_evidence_registered_at=registered_at)
            orders.append(order)
        for order in orders[2:4]:
            order.refresh_from_db()
            order.customer_name_snapshot = "Synthetic Payment Customer"
            order.save(update_fields=("customer_name_snapshot",))
        outside_order = self.preview_and_confirm(self.order_payload(), actor=self.sales_two)[2]
        outside_order.refresh_from_db()
        outside_image = self.add_payment_image(owner=self.sales_two)
        register_payment(self.sales_two, outside_order.pk, {
            "operation_key": str(uuid4()),
            "expected_version": outside_order.version,
            "account_id": self.public_account.pk,
            "received_amount": str(outside_order.revenue_amount),
            "image_ids": [outside_image.pk],
            "difference_reason": "",
        })
        PaymentRecord.objects.filter(order=outside_order).update(
            first_evidence_registered_at=registration_times[2],
        )
        outside_order.refresh_from_db()
        outside_order.customer_name_snapshot = "Synthetic Payment Customer"
        outside_order.save(update_fields=("customer_name_snapshot",))

        fixed_now = datetime(2026, 10, 31, 12, 0, tzinfo=SHANGHAI)
        custom_query = {"period": "custom", "start_date": "2026-10-01", "end_date": "2026-10-01"}
        with patch("crm.reports._now_local", return_value=fixed_now):
            rolling = build_report(self.sales, "payments", {"period": "1m"})
            custom = build_report(self.sales, "payments", custom_query)
            searched = build_report(self.sales, "payments", {
                **custom_query,
                "q": "Synthetic Payment",
                "account_id": str(self.public_account.pk),
                "status": PaymentRecord.Status.PENDING,
            })
            searched_by_id = build_report(self.sales, "payments", {
                **custom_query,
                "q": str(orders[2].pk),
            })

        self.assertEqual((rolling["metadata"]["start_date"], rolling["metadata"]["end_date"]), ("2026-09-30", "2026-10-31"))
        self.assertEqual({row["order_id"] for row in rolling["items"]}, {str(order.pk) for order in orders[1:]})
        self.assertEqual({row["order_id"] for row in custom["items"]}, {str(orders[2].pk), str(orders[3].pk)})
        self.assertEqual({row["order_id"] for row in searched["items"]}, {str(orders[2].pk), str(orders[3].pk)})
        self.assertNotIn(str(outside_order.pk), {row["order_id"] for row in searched["items"]})
        self.assertEqual([row["order_id"] for row in searched_by_id["items"]], [str(orders[2].pk)])
        self.assertEqual(searched["items"][0]["order_status"], orders[2].status)
        self.assertEqual(searched["items"][0]["order_status_label"], orders[2].get_status_display())
        self.assertEqual(searched["items"][0]["status"], PaymentRecord.Status.PENDING)

        self.client.force_login(self.sales)
        query = (
            "period=custom&start_date=2026-10-01&end_date=2026-10-01"
            f"&q=Synthetic+Payment&account_id={self.public_account.pk}&status={PaymentRecord.Status.PENDING}"
        )
        with patch("crm.reports._now_local", return_value=fixed_now):
            api = self.client.get(f"/api/v1/reports/payments?{query}&page_size=200")
            page = self.client.get(f"/reports/payments/?{query}&page=2&page_size=1")
            export = self.client.get(f"/api/v1/reports/payments/export?{query}&page=2&page_size=1")

        self.assertEqual(api.status_code, 200)
        self.assertEqual(api.json()["total"], 2)
        self.assertEqual({row["order_id"] for row in api.json()["items"]}, {str(orders[2].pk), str(orders[3].pk)})
        self.assertEqual(
            {column["label"] for column in api.json()["columns"] if column["key"] in ("order_status_label", "status_label")},
            {"履约状态", "核对状态"},
        )
        self.assertTrue(all(row["order_status_label"] for row in api.json()["items"]))
        self.assertEqual(page.status_code, 200)
        self.assertEqual(page.context["total"], 2)
        self.assertEqual(page.context["page"], 2)
        self.assertContains(page, "首次登记收款凭证时间")
        self.assertEqual(export.status_code, 200)
        from io import BytesIO
        from openpyxl import load_workbook

        sheet = load_workbook(BytesIO(export.content), data_only=False).active
        self.assertEqual(sheet.max_row, 8)
        headers = [sheet.cell(6, column).value for column in range(1, sheet.max_column + 1)]
        order_id_column = headers.index("订单编号") + 1
        exported_ids = {str(sheet.cell(row, order_id_column).value) for row in range(7, 9)}
        self.assertEqual(exported_ids, {str(orders[2].pk), str(orders[3].pk)})
        self.assertIn("履约状态", headers)
        self.assertIn("核对状态", headers)
        registered_at_column = headers.index("首次登记收款凭证时间") + 1
        exported_dates = [sheet.cell(row, registered_at_column).value for row in range(7, 9)]
        api_dates = [row["first_evidence_registered_at"] for row in api.json()["items"]]
        self.assertEqual(exported_dates, api_dates)
        self.assertEqual(export["X-Report-Start-Date"], "2026-10-01")
        self.assertEqual(export["X-Report-End-Date"], "2026-10-01")

    def test_commission_uses_original_earned_date_and_lists_deltas_separately(self):
        order = self.new_order()
        self.register(order)
        self.verify(order)
        order.refresh_from_db()
        deliver_on_site(self.user, order.pk, {"operation_key": "8271cae4-2a34-4a55-9c15-7e64c9d0ae2c", "expected_version": order.version})
        order.refresh_from_db()
        earned_at = datetime(2025, 9, 30, 16, 30, tzinfo=ZoneInfo("UTC"))
        CommissionReview.objects.create(
            order=order, approved_amount=Decimal("1.000"), approved_amount_version=order.amount_version,
            first_earned_at=earned_at, reviewer=self.owner, reviewed_at=earned_at,
        )
        operations = list(BusinessOperation.objects.filter(sales_order=order).order_by("sequence"))
        self.assertGreaterEqual(len(operations), 3)
        CommissionAdjustment.objects.create(
            approval_operation=operations[0], order=order, kind=CommissionAdjustment.Kind.INITIAL,
            old_amount=Decimal("0.000"), new_amount=Decimal("1.000"), delta=Decimal("1.000"), confirmed_at=earned_at,
        )
        CommissionAdjustment.objects.create(
            approval_operation=operations[1], order=order, kind=CommissionAdjustment.Kind.REREVIEW,
            old_amount=Decimal("1.000"), new_amount=Decimal("1.200"), delta=Decimal("0.200"),
            confirmed_at=datetime(2025, 10, 2, 1, 0, tzinfo=SHANGHAI), completed_by=self.sales,
            completed_at=datetime(2025, 10, 2, 2, 0, tzinfo=SHANGHAI),
        )
        CommissionAdjustment.objects.create(
            approval_operation=operations[2], order=order, kind=CommissionAdjustment.Kind.COST_CORRECTION,
            old_amount=Decimal("1.200"), new_amount=Decimal("0.800"), delta=Decimal("-0.400"),
            confirmed_at=datetime(2025, 10, 3, 1, 0, tzinfo=SHANGHAI),
        )

        earned = build_report(self.sales, "commissions", {
            "period": "custom", "start_date": "2025-10-01", "end_date": "2025-10-01",
        })
        self.assertEqual([(row["entry_type"], row["approved_amount"]) for row in earned["items"]], [("order", "1.000")])
        SalesOrder.objects.filter(pk=order.pk).update(amount_version=order.amount_version + 1)
        needs_review = build_report(self.sales, "commissions", {
            "period": "custom", "start_date": "2025-10-01", "end_date": "2025-10-01",
        })["items"][0]
        self.assertEqual(needs_review["status"], "needs_review")
        self.assertEqual(needs_review["approved_amount"], "1.000")
        self.assertTrue(needs_review["review_delta"])
        rereview = build_report(self.sales, "commissions", {
            "period": "custom", "start_date": "2025-10-02", "end_date": "2025-10-02",
        })
        self.assertEqual(len(rereview["items"]), 1)
        self.assertEqual(rereview["items"][0]["adjustment_kind"], "重审差额")
        self.assertEqual(rereview["items"][0]["delta_amount"], "0.200")
        self.assertEqual(rereview["items"][0]["completion_status"], "已完成")
        correction = build_report(self.sales, "commissions", {
            "period": "custom", "start_date": "2025-10-03", "end_date": "2025-10-03",
        })
        self.assertEqual(correction["items"][0]["delta_amount"], "-0.400")
        self.assertEqual(correction["summary"]["adjustment_delta_not_added"], "-0.400")
        self.assertEqual(correction["summary"]["states"], {})

    def test_commission_visibility_requires_its_own_permission(self):
        self.new_order()
        _, _, other_order = self.preview_and_confirm(actor=self.sales_two)
        self.client.force_login(self.sales)
        response = self.client.get("/api/v1/reports/commissions")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["total"], 1)

        view_all_orders = Permission.objects.get(content_type__app_label="crm", codename="view_all_sales_orders")
        view_all_commissions = Permission.objects.get(content_type__app_label="crm", codename="view_all_commissions")
        self.sales.user_permissions.add(view_all_commissions)
        self.sales = self.User.objects.get(pk=self.sales.pk)
        self.client.force_login(self.sales)
        response = self.client.get("/api/v1/reports/commissions")
        self.assertEqual(response.json()["total"], 1)
        self.assertNotIn(str(other_order.pk), [row["order_id"] for row in response.json()["items"]])

        self.sales.user_permissions.add(view_all_orders)
        self.sales = self.User.objects.get(pk=self.sales.pk)
        self.client.force_login(self.sales)
        response = self.client.get("/api/v1/reports/commissions")
        self.assertEqual(response.json()["total"], 2)
        self.assertIn(str(other_order.pk), [row["order_id"] for row in response.json()["items"]])

        self.sales.user_permissions.remove(view_all_commissions)
        self.sales = self.User.objects.get(pk=self.sales.pk)
        self.client.force_login(self.sales)
        response = self.client.get("/api/v1/reports/commissions")
        self.assertEqual(response.json()["total"], 1)
        self.assertNotIn(str(other_order.pk), [row["order_id"] for row in response.json()["items"]])

    def test_export_keeps_all_rows_text_numbers_and_negative_amount(self):
        self.balance.sku.barcode = "0012345678901"
        self.balance.sku.save(update_fields=("barcode",))
        second_warehouse = Warehouse.objects.create(name="Synthetic export warehouse", address="Synthetic address")
        InventoryBalance.objects.create(sku=self.sku, warehouse=second_warehouse, on_hand=1, cost_amount=Decimal("1.235"))
        self.client.force_login(self.sales)
        page = self.client.get("/api/v1/reports/inventory?page_size=1")
        self.assertEqual(page.status_code, 200, page.content.decode())
        self.assertEqual(page.json()["total"], 2)
        export = self.client.get("/api/v1/reports/inventory/export")
        self.assertEqual(export.status_code, 200)
        from io import BytesIO
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(export.content), data_only=False)
        sheet = workbook.active
        self.assertEqual(sheet.max_row, 8)
        self.assertEqual(sheet["E7"].value, "0012345678901")
        self.assertEqual(sheet["E7"].data_type, "s")
        data = {
            "title": "测试导出",
            "metadata": {
                "period_label": "全部日期", "date_basis": "合成验证", "filter_summary": "=1+1",
                "queried_at": "2026-10-05T12:00:00+08:00",
            },
            "columns": [("id", "编号", "text"), ("amount", "差额", "amount")],
            "items": [{"id": "0012", "amount": "-1.235"}],
        }
        workbook = load_workbook(BytesIO(workbook_bytes(data)), data_only=False)
        sheet = workbook.active
        self.assertEqual(sheet["A7"].value, "0012")
        self.assertEqual(sheet["A7"].data_type, "s")
        self.assertEqual(sheet["B7"].value, -1.235)
        self.assertEqual(sheet["B7"].data_type, "n")
        self.assertIn(".000", sheet["B7"].number_format)
        self.assertEqual(sheet["B4"].value, "=1+1")
        self.assertEqual(sheet["B4"].data_type, "s")
        data["items"] = [{"id": "0012", "amount": "123456789012345.678"}]
        with self.assertRaises(OperationError) as raised:
            workbook_bytes(data)
        self.assertEqual(raised.exception.code, "export_amount_precision")
